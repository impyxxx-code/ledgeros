#!/usr/bin/env python3
"""Build the ARKHAM RETAIL prospect workbook (XLSX) from output/master.csv.

Sheets:
  DASHBOARD (live COUNTIF formulas), ALL PROSPECTS, VAPE SHOPS, CONVENIENCE,
  PETROL - FORECOURTS, OFF LICENCES, TOBACCO - SMOKE - SHISHA, SUPERMARKETS,
  CHAINS - MULTI-SITE, MISSING CONTACT INFO, DUPLICATE REVIEW, SOURCE LOG,
  CRM - FOLLOW-UP
"""

import csv
import json
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parents[1]
TODAY = date.today().isoformat()

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name="Arial", size=10)
TITLE_FONT = Font(name="Arial", bold=True, size=14, color="1F3864")
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ALT_FILL = PatternFill("solid", fgColor="F2F6FC")

with open(BASE / "output" / "master.csv", encoding="utf-8-sig") as f:
    reader = csv.DictReader(f)
    COLUMNS = reader.fieldnames
    ROWS = list(reader)

with open(BASE / "data" / "master.json") as f:
    RECORDS = json.load(f)

COL_IDX = {c: i + 1 for i, c in enumerate(COLUMNS)}


def col_letter(name, columns=None):
    columns = columns or COLUMNS
    return get_column_letter(columns.index(name) + 1)


def write_table(ws, columns, rows, freeze="A2", widths=None):
    for j, c in enumerate(columns, 1):
        cell = ws.cell(row=1, column=j, value=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=False)
        cell.border = BORDER
    for i, row in enumerate(rows, 2):
        for j, c in enumerate(columns, 1):
            cell = ws.cell(row=i, column=j, value=row.get(c, ""))
            cell.font = BODY_FONT
            cell.border = BORDER
            if i % 2 == 0:
                cell.fill = ALT_FILL
    ws.freeze_panes = freeze
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    default_w = 18
    for j, c in enumerate(columns, 1):
        w = (widths or {}).get(c)
        if w is None:
            longest = max([len(str(c))] + [len(str(r.get(c, ""))) for r in rows[:200]])
            w = min(max(longest + 2, 10), 45)
        ws.column_dimensions[get_column_letter(j)].width = w or default_w


KEY_COLS = [
    "Record ID", "Location ID", "Parent Company", "Business Name",
    "Business Category", "Business Subcategory", "Independent / Chain",
    "Number of Locations", "Address Line 1", "Town / City", "Postcode",
    "Main Phone", "General Email", "Sales Email", "Website", "Facebook",
    "Instagram", "Contact Name", "Contact Position",
    "Companies House Number", "Company Status", "Lead Priority",
    "Priority Reasons", "Phone Status", "Email Status", "Data Confidence",
    "Verification Status", "Primary Source",
]

wb = Workbook()

# ---------------- DASHBOARD ----------------
dash = wb.active
dash.title = "DASHBOARD"
dash.sheet_view.showGridLines = False
dash["B2"] = "ARKHAM RETAIL - UK RETAIL PROSPECT DATABASE"
dash["B2"].font = TITLE_FONT
dash["B3"] = f"Pilot region: West Yorkshire (Yorkshire & Humber) - generated {TODAY}"
dash["B3"].font = Font(name="Arial", italic=True, size=10, color="666666")

n = len(ROWS)
last = n + 1  # last data row in ALL PROSPECTS
AP = "'ALL PROSPECTS'"


def rng(colname):
    L = col_letter(colname)
    return f"{AP}!${L}$2:${L}${last}"


# each metric: (label, live formula, python snapshot of the same criterion)
metrics = [
    ("TOTAL PROSPECTS", f"=COUNTA({rng('Record ID')})",
     sum(1 for x in ROWS if x["Record ID"])),
    ("PHONE AVAILABLE", f"=COUNTIF({rng('Main Phone')},\"?*\")",
     sum(1 for x in ROWS if x["Main Phone"])),
    ("EMAIL AVAILABLE", f"=SUMPRODUCT(--(({rng('General Email')}<>\"\")+({rng('Sales Email')}<>\"\")+({rng('Purchasing Email')}<>\"\")>0))",
     sum(1 for x in ROWS if x["General Email"] or x["Sales Email"] or x["Purchasing Email"])),
    ("WEBSITE AVAILABLE", f"=COUNTIF({rng('Website')},\"?*\")",
     sum(1 for x in ROWS if x["Website"])),
    ("PHONE VERIFIED (OFFICIAL SITE)", f"=COUNTIF({rng('Phone Status')},\"Verified on official website\")",
     sum(1 for x in ROWS if x["Phone Status"] == "Verified on official website")),
    ("EMAIL PUBLICLY LISTED", f"=COUNTIF({rng('Email Status')},\"Publicly Listed\")+COUNTIF({rng('Email Status')},\"Verified\")",
     sum(1 for x in ROWS if x["Email Status"] in ("Publicly Listed", "Verified"))),
    ("COMPANIES HOUSE VERIFIED", f"=COUNTIF({rng('Companies House Checked')},\"Yes\")",
     sum(1 for x in ROWS if x["Companies House Checked"] == "Yes")),
    ("REQUIRES VERIFICATION", f"=COUNTIF({rng('Verification Status')},\"Requires Verification\")",
     sum(1 for x in ROWS if x["Verification Status"] == "Requires Verification")),
    ("DUPLICATE REVIEW REQUIRED", f"=COUNTIF({rng('Duplicate Review')},\"?*\")",
     sum(1 for x in ROWS if x["Duplicate Review"])),
    ("MULTI-SITE LOCATIONS", f"=COUNTIF({rng('Parent Company')},\"?*\")",
     sum(1 for x in ROWS if x["Parent Company"])),
    ("HIGH PRIORITY LEADS", f"=COUNTIF({rng('Lead Priority')},\"High\")",
     sum(1 for x in ROWS if x["Lead Priority"] == "High")),
    ("MEDIUM PRIORITY LEADS", f"=COUNTIF({rng('Lead Priority')},\"Medium\")",
     sum(1 for x in ROWS if x["Lead Priority"] == "Medium")),
]
r = 5
dash.cell(row=r, column=2, value="KEY METRICS").font = Font(name="Arial", bold=True, size=11)
dash.cell(row=r, column=3, value="LIVE").font = Font(name="Arial", bold=True, size=9, color="808080")
dash.cell(row=r, column=4, value=f"SNAPSHOT {TODAY}").font = Font(name="Arial", bold=True, size=9, color="808080")
for i, (label, formula, snap) in enumerate(metrics):
    rr = r + 1 + i
    lc = dash.cell(row=rr, column=2, value=label)
    vc = dash.cell(row=rr, column=3, value=formula)
    sc = dash.cell(row=rr, column=4, value=snap)
    lc.font = BODY_FONT
    vc.font = Font(name="Arial", bold=True, size=10)
    sc.font = Font(name="Arial", size=10, color="808080")
    lc.border = BORDER
    vc.border = BORDER
    sc.border = BORDER
    vc.number_format = sc.number_format = "#,##0"
    if i % 2 == 0:
        lc.fill = ALT_FILL
        vc.fill = ALT_FILL
        sc.fill = ALT_FILL

cats = ["Vape Shop", "Convenience Store", "Petrol Station / Forecourt", "Off-Licence",
        "Tobacco Specialist", "Newsagent", "Independent Supermarket",
        "Smoke / Shisha Shop", "Cash & Carry", "Other Relevant Retail"]
dash.cell(row=r, column=6, value="BY CATEGORY").font = Font(name="Arial", bold=True, size=11)
for i, cat in enumerate(cats):
    rr = r + 1 + i
    lc = dash.cell(row=rr, column=6, value=cat)
    vc = dash.cell(row=rr, column=7, value=f"=COUNTIF({rng('Business Category')},\"{cat}\")")
    sc = dash.cell(row=rr, column=8, value=sum(1 for x in ROWS if x["Business Category"] == cat))
    lc.font, vc.font = BODY_FONT, Font(name="Arial", bold=True, size=10)
    sc.font = Font(name="Arial", size=10, color="808080")
    lc.border = vc.border = sc.border = BORDER
    if i % 2 == 0:
        lc.fill = vc.fill = sc.fill = ALT_FILL

towns = sorted({row["Town / City"] for row in ROWS if row["Town / City"]})
dash.cell(row=r, column=10, value="BY TOWN / CITY").font = Font(name="Arial", bold=True, size=11)
for i, t in enumerate(towns[:34]):
    rr = r + 1 + i
    lc = dash.cell(row=rr, column=10, value=t)
    vc = dash.cell(row=rr, column=11, value=f"=COUNTIF({rng('Town / City')},\"{t}\")")
    sc = dash.cell(row=rr, column=12, value=sum(1 for x in ROWS if x["Town / City"] == t))
    lc.font, vc.font = BODY_FONT, Font(name="Arial", bold=True, size=10)
    sc.font = Font(name="Arial", size=10, color="808080")
    lc.border = vc.border = sc.border = BORDER
    if i % 2 == 0:
        lc.fill = vc.fill = sc.fill = ALT_FILL

dash.cell(row=r + 14, column=2,
          value="LIVE column recalculates from ALL PROSPECTS when opened in Excel/Sheets; "
                "SNAPSHOT is the static count at generation time. If they differ, the data was edited."
          ).font = Font(name="Arial", italic=True, size=9, color="808080")

for col, w in {"A": 3, "B": 34, "C": 10, "D": 14, "E": 3, "F": 28, "G": 8, "H": 12,
               "I": 3, "J": 24, "K": 8, "L": 12}.items():
    dash.column_dimensions[col].width = w

# ---------------- ALL PROSPECTS ----------------
ws = wb.create_sheet("ALL PROSPECTS")
write_table(ws, COLUMNS, ROWS)

# ---------------- Category sheets ----------------
def subset(pred):
    return [row for row in ROWS if pred(row)]

sheets = [
    ("VAPE SHOPS", lambda r: r["Business Category"] == "Vape Shop"),
    ("CONVENIENCE", lambda r: r["Business Category"] in ("Convenience Store", "Newsagent")),
    ("PETROL - FORECOURTS", lambda r: r["Business Category"] == "Petrol Station / Forecourt"),
    ("OFF LICENCES", lambda r: r["Business Category"] == "Off-Licence"),
    ("TOBACCO - SMOKE - SHISHA", lambda r: r["Business Category"] in ("Tobacco Specialist", "Smoke / Shisha Shop")),
    ("SUPERMARKETS", lambda r: r["Business Category"] in ("Independent Supermarket", "Cash & Carry")),
    ("CHAINS - MULTI-SITE", lambda r: bool(r["Parent Company"]) or r["Independent / Chain"] in
        ("Multi-Site Independent", "Regional Chain", "National Chain")),
    ("MISSING CONTACT INFO", lambda r: not r["Main Phone"] and not (r["General Email"] or r["Sales Email"] or r["Purchasing Email"])),
    ("DUPLICATE REVIEW", lambda r: bool(r["Duplicate Review"])),
]
for name, pred in sheets:
    ws = wb.create_sheet(name)
    write_table(ws, KEY_COLS, subset(pred))

# ---------------- SOURCE LOG ----------------
src_cols = ["Record ID", "Business Name", "Town / City", "Source URL", "Source Type",
            "Date Collected", "Information Obtained", "Verification Status"]
src_rows = []
rec_by_id = {r["record_id"]: r for r in RECORDS}
for row in ROWS:
    rec = rec_by_id.get(row["Record ID"], {})
    for s in rec.get("sources", []):
        src_rows.append({
            "Record ID": row["Record ID"], "Business Name": row["Business Name"],
            "Town / City": row["Town / City"], "Source URL": s.get("url", ""),
            "Source Type": s.get("type", ""), "Date Collected": s.get("date", ""),
            "Information Obtained": s.get("info_obtained", ""),
            "Verification Status": row["Verification Status"],
        })
ws = wb.create_sheet("SOURCE LOG")
write_table(ws, src_cols, src_rows)

# ---------------- CRM ----------------
crm_cols = ["Record ID", "Business Name", "Town / City", "Business Category",
            "Lead Priority", "Main Phone", "General Email", "Sales Email",
            "Contact Name", "Contact Position", "Lead Status", "Date Added",
            "Last Contact", "Next Follow-Up", "Contact Method", "Salesperson",
            "Outcome", "CRM Notes", "Opt-Out Status", "Do Not Contact",
            "Marketing Legal Basis"]
ws = wb.create_sheet("CRM - FOLLOW-UP")
prio_order = {"High": 0, "Medium": 1, "Low": 2, "Unclassified": 3}
crm_rows = sorted(ROWS, key=lambda r: (prio_order.get(r["Lead Priority"], 4), r["Business Name"]))
write_table(ws, crm_cols, crm_rows)
# example row legend
note = ws.cell(row=len(crm_rows) + 3, column=1,
               value="Editable CRM columns: Lead Status, Last Contact, Next Follow-Up, Contact Method, "
                     "Salesperson, Outcome, CRM Notes, Opt-Out Status, Do Not Contact. "
                     "Never delete a row with Opt-Out/Do Not Contact set - suppression must survive re-imports.")
note.font = Font(name="Arial", italic=True, size=9, color="808080")

out = BASE / "output" / "ARKHAM_RETAIL_prospects_west_yorkshire.xlsx"
wb.save(out)
print(f"Saved {out} - {len(ROWS)} prospects, {len(src_rows)} source-log rows, {len(wb.sheetnames)} sheets")
